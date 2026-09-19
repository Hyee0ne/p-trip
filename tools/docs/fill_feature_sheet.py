#!/usr/bin/env python3
"""공모전 「기능설명서」 양식(pptx)을 P의 여행 내용으로 채운다.

입력: docs/2026 … 기능설명서 양식(작성용).pptx (원본은 손대지 않는다)
출력: docs/기능설명서_P의여행.pptx
이미지: docs/screenshots/*.png + IMG_DIR(시뮬레이터 추가 캡처). 없는 파일은 건너뛰고 알린다.

실행: python3 tools/docs/fill_feature_sheet.py [IMG_DIR]
"""
import copy, pathlib, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# ⚠ 표의 첫 행은 템플릿 표 스타일(firstRow)이 흰 글자다 — 색을 안 주면 흰 칸에 흰 글자로 사라진다.
#   (원본 가이드 문구는 빨간색을 명시해서 보였다.) 내용 글자는 전부 잉크색을 명시한다.
INK = RGBColor(0x1F, 0x23, 0x26)

ROOT = pathlib.Path(__file__).resolve().parents[2]
SRC = ROOT / 'docs' / '2026 관광데이터 활용 공모전 웹앱 개발 부문 기능설명서 양식(작성용).pptx'
OUT = ROOT / 'docs' / '기능설명서_P의여행.pptx'
SHOTS = ROOT / 'docs' / 'screenshots'
IMG_DIR = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else None
ICON = ROOT / 'docs/썸네일/썸네일_밝음.png'  # 대표 이미지(썸네일). 아이콘 대신 로고+이름 이미지 (2026-09-16)
SHARE = ROOT / 'app/test/goldens/share_card_8stops.png'

missing = []
import json as _json
_meta_p = SHOTS / 'case' / 'meta.json'
CASE = _json.loads(_meta_p.read_text()) if _meta_p.exists() else {'date': '', 'market_day': False}
CASE_LABEL = (f"{CASE['date']} 장날 실제 캡처" if CASE.get('market_day') else f"{CASE['date']} 비장날 캡처 (장날 캡처로 교체 예정)") if CASE.get('date') else '재현 캡처'
CARD_LINE = '「오늘이 마침 북평민속오일장이에요」' if CASE.get('market_day') else '「다음 장은 N일 뒤 · 여기서 약 N km」 (비장날엔 다음 일정을 안내)'

def img(name):
    """스크린샷 이름 → 경로. 스토어 스크린샷 → 추가 캡처 → 골든 순."""
    for base in (SHOTS, SHOTS / 'case', IMG_DIR, ROOT / 'app/test/goldens'):
        if base is None: continue
        p = base / name
        if p.exists(): return p
    missing.append(name); return None

def set_cell(cell, lines, size=13, bold_first=False, align=PP_ALIGN.LEFT, color=INK, gap=5):
    """셀 글을 통째로 바꾼다. lines: str 또는 [str]. 빈 줄은 '' 로.
    줄바꿈 규칙 (2026-09-16, 가독성): '• 제목' / '1. 제목' 줄은 **굵게** 하고 앞에 간격을 둔다.
    '- 세부' 줄은 들여쓰기(≒0.22in)하고 '‒ ' 로 시작한다. 한 줄에 한 생각."""
    from pptx.oxml.ns import qn
    import re as _re
    if isinstance(lines, str): lines = [lines]
    tf = cell.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs): r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        p = first if i == 0 else tf.add_paragraph()
        p.alignment = align
        head = bool(_re.match(r'^(• |\d+\. |[①②③④⑤])', line))
        sub = line.startswith('- ')
        text = ('‒ ' + line[2:]) if sub else line
        if sub:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(Inches(0.22)))); pPr.set('indent', str(-int(Inches(0.14))))
        if head and i > 0: p.space_before = Pt(gap)
        run = p.add_run(); run.text = text
        run.font.size = Pt(size if not sub else size - 0.5)
        run.font.bold = True if ((bold_first and i == 0) or head) else False
        if color is not None: run.font.color.rgb = color
        if sub: run.font.color.rgb = RGBColor(0x3E, 0x43, 0x47)

def find_table(slide, name=None):
    for sh in slide.shapes:
        if sh.has_table and (name is None or sh.name == name): return sh
    raise KeyError(name)

def remove_shape(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            sh._element.getparent().remove(sh._element); return True
    return False

def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sldId = list(sldIdLst)[index]
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

def duplicate_slide(prs, index):
    """그림 없는 슬라이드(표·도형만)를 복제해 맨 뒤에 붙인다. 위치는 호출부가 옮긴다."""
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):  # 레이아웃이 만든 자리표시자 제거
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        # ⚠ 그림은 복사하지 않는다 (2026-09-16). 원본에 이미 넣은 캡처까지 딸려 와 새 슬라이드의 그림 아래에
        #   같은 이미지가 다른 크기로 깔렸다 — 셀마다 두 겹이 겹쳐 '배경이 비치는' 것처럼 보였다.
        if sh.shape_type == 13:
            continue
        dst.shapes._spTree.append(copy.deepcopy(sh._element))
    return dst

def move_slide(prs, old_index, new_index):
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    el = ids[old_index]
    sldIdLst.remove(el)
    sldIdLst.insert(new_index, el)

def add_shot(slide, path, left, top, box_w, box_h, pad=0.06, crop=None):
    """상자(인치) 안에 비율 유지로 가운데 놓는다. crop=(top, bottom) 은 세로 비율로 잘라 핵심만 키운다."""
    from PIL import Image
    import tempfile
    if crop:
        with Image.open(path) as im:
            w0, h0 = im.size
            box = (0, int(h0 * crop[0]), w0, int(h0 * crop[1]))
            tmp = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
            im.crop(box).save(tmp.name); path = pathlib.Path(tmp.name)
    with Image.open(path) as im: w, h = im.size
    bw, bh = box_w - pad * 2, box_h - pad * 2
    scale = min(bw / w, bh / h)
    pw, ph = w * scale, h * scale
    slide.shapes.add_picture(str(path), Inches(left + pad + (bw - pw) / 2), Inches(top + pad + (bh - ph) / 2), Inches(pw), Inches(ph))

# ──────────────────────────────── 내용 ────────────────────────────────
TEAM = '라이스쿠키'
NAME = 'P의여행:국도편'  # App Store 이름이 메인
APP_STORE_URL = 'https://apps.apple.com/kr/app/id6806749553'
VIDEO_URL = 'https://youtu.be/q4jvO63vvQA'  # 데모 영상. 비어 있으면 줄을 넣지 않는다.

INTRO = {
    'name': 'P의여행:국도편  (국도 위의 발견 레이더)',
    'type': '앱 서비스 (iOS) · App Store 출시',
    'summary': [
        '목적지는 정하지 않아도 됩니다. 달릴 국도와 방향만 고르세요.',
        'P의여행은 지금 가는 길 앞에서 들를 만한 곳을 알려줍니다. 장날·일몰 정보를 더해 "왜 지금 들르면 좋은지"를 설명하고, 선택한 장소는 티맵·애플 지도로 안내합니다.',
        '길안내는 내비가, 뜻밖의 발견은 P의여행이.',
        '이용 대상: 목적지 없이 국도 드라이브를 떠나는 운전자와 동승자',
    ],
    'why': [
        '• 「풍향중」처럼, 계획 없이 떠나는 여행을 일상에서도',
        '- 즉흥 여행은 뜻밖의 발견이 즐겁습니다. 하지만 막상 떠나면 어디에 들를지, 오늘 장이 서는지 다시 찾아야 합니다.',
        '- P의여행은 지금 가는 길에서 들를 이유를 먼저 알려줍니다.',
        '• 기대효과와 검증',
        '- 목적지까지 통과하던 여행자가 계획에 없던 지역의 장소에 한 번 더 멈추게 합니다.',
        '- 정차 감지로 방문 후보를 기록하고, 여행 후 사용자 확인으로 실제 방문과 계획에 없던 방문 여부를 검증합니다. (기기 내 집계·동의 기반)',
    ],
}

FEATURES_LIST = [
    '1. 국도 선택 지도 — 목적지 대신, 달릴 길을 고릅니다',
    '- 지도 위 국도 51선 가운데 지금 여기서 탈 수 있는 길을 보여줍니다. 길마다 오늘 장이 서는 곳과 함께 찾는 명소를 한 줄로 소개합니다.',
    '2. 발견 레이더 — 지금 들를 만한 곳을 알려줍니다',
    '- 현재 위치와 진행 방향을 기준으로 반경 5km 안의 장소를 카드 한 장씩 보여줍니다. 길을 바꿔도 그대로 작동합니다.',
    '3. 오늘만 볼 수 있는 발견 — "왜 지금 들르면 좋은지"를 설명합니다',
    '- 장날·일몰·이번 주에 끝나는 행사를 관광정보에 더해 먼저 보여줍니다. 장이 서지 않는 날에는 다음 장날을 알려줍니다.',
    '4. 소리·알림과 내비 연결 — 발견은 듣고, 길안내는 내비로',
    '- 이동 중에는 음성으로 안내하고, 놓친 발견은 알림 센터에 남깁니다. 「들르기」 한 번으로 티맵·애플 지도가 그곳까지 안내합니다.',
    '- 쉬었다가 다시 떠날 때는 앞쪽 10km 안의 다음 들를 곳을 최대 3곳 제안하고, 다른 국도로 갈아타도 기록이 이어집니다.',
    '5. 여행기 — 지나온 길이 여행기로 남습니다',
    '- 하루가 「N번 국도에서 생긴 일」 한 편으로 자동 기록됩니다. GPS 포스터·타임라인·공유 카드, 국도 51선 수집. 계정 없이 기기 안에 저장합니다.',
    '',
    '• 추천 기준 · 이용 안내',
    '- 사진이 있고 영업시간·전화·주소·소개 같은 확인 가능한 정보를 일정 기준(신뢰도 60점) 이상 갖춘 장소만 추천합니다. 표시 거리는 직선거리이며, 장날은 안내하되 임시휴무·영업시간 변동은 방문 전 확인이 필요합니다. 장소 선택은 동승자 또는 정차한 운전자가 진행합니다.',
]

# 핵심 기능별 흐름도 (최대 5장). 캡처 4장 + 단계 설명 4칸.
FLOWS = [
    {
        'title': '출발 — 목적지 대신, 달릴 길을 고릅니다',
        'desc': ['지도에서 지금 여기서 탈 수 있는 국도를 고르고, 방향만 정하면 출발입니다.',
                 '앱은 그 국도의 진입 지점을 티맵·애플 지도에 넘겨 길안내를 맡깁니다. 목적지는 입력하지 않습니다.'],
        'shots': ['1_home.png', 'home_expanded.png', '2_depart.png', 'handoff-t16.png'],
        'steps': [
            ['① 홈 지도', '- 국도 51선과 「여기서 탈 수 있는 길」'],
            ['② 51선 전체 목록', '- 남북(홀수)·동서(짝수) 필터'],
            ['③ 방향 선택', '- 「여기서 어느 쪽으로 갈까요?」 북/남'],
            ['④ 내비 연결', '- 진입 지점까지 티맵 / 애플 지도가 안내'],
        ],
    },
    {
        'title': '발견 — 지금 들를 만한 곳을 알려줍니다',
        'desc': ['현재 위치와 진행 방향을 기준으로 주변 장소를 추천합니다. 이동 중에는 음성으로 안내하고, 놓친 발견은 알림 센터에 남깁니다.',
                 '추천 범위는 반경 5km이며, 표시 거리는 직선거리입니다. 대표 사례: 7번 국도 · 북평민속오일장(3·8일 장)'],
        'shots': ['4_radar.png', ('case-card.png', (0.34, 0.785)), 'notif-lock.png', 'notif-sunset.png'],
        'steps': [
            ['① 진행 방향의 장소 탐색', '- 레이더: 내 위치 기준 앞쪽 스팟'],
            ['② 장날·거리 안내 카드', f'- {CARD_LINE}', f'- {CASE_LABEL}'],
            ['③ 잠금화면 알림 (확대)', '- 「오늘이 마침 경안시장이에요」 (2026-09-13 실주행)', '- 카드 문장은 음성으로도 낭독(내비 음성을 낮췄다 되돌림)'],
            ['④ 내비 사용 중 알림 배너', '- 티맵 위 「곧 하남 나무고아원에 해가 져요」 (실주행)'],
        ],
    },
    {
        'title': '방문 — 마음에 들면, 내비로 연결합니다',
        'desc': ['「들르기」 한 번으로 티맵·애플 지도가 그곳까지 안내합니다. 앱은 길안내를 직접 하지 않습니다.',
                 '들른 곳은 「오늘 들른 곳」 자취로 남습니다. 장소 선택은 동승자 또는 정차한 운전자가 진행합니다.'],
        'shots': [('case-card.png', (0.45, 1.0)), ('case-handoff.png', (0.4, 1.0)), 'case-spot.png', 'case-trail.png'],
        'steps': [
            ['① 들르기 / 찜', '- 카드 아래 넘기기 · 들르기 · 찜'],
            ['② 내비 연결', '- 「북평민속오일장」 → 티맵 / 애플 지도', f'- {CASE_LABEL}'],
            ['③ 장소 상세', '- 사진 · 소개 · 영업시간 · 장날', '- 별점 없이 「네이버 후기 보기」 링크만'],
            ['④ 오늘 들른 곳', '- 레이더 아래 자취에 사진으로 기록'],
        ],
    },
    {
        'title': '재출발 — 쉬었다가, 다음 들를 곳을 찾습니다',
        'desc': ['쉬었다가 다시 떠날 때, 다음 들를 곳을 제안합니다. 2분간 정차하면 진행 방향 10km 안의 후보를 최대 3곳 보여줍니다.',
                 '다른 국도로 갈아타도 여행과 기록은 이어집니다. 달리는 중에는 열리지 않습니다.'],
        'shots': ['next2-t17.png', 'next2-t22.png', 'switch-t16.png', 'radar-t14.png'],
        'steps': [
            ['① 정차 2분 → 자동 표시', '- 「여기서 앞쪽으로」'],
            ['② 앞쪽 후보', '- 10km 안 최대 3곳 · 「그냥 7번 국도로 돌아가기」'],
            ['③ 길 바꾸기', '- 「길을 바꿀까요?」 인근 국도로 전환'],
            ['④ 기록 유지', '- 여행은 하루, 국도는 구간(43 → 6번)'],
        ],
    },
    {
        'title': '기록 — 지나온 길이 여행기로 남습니다',
        'desc': ['「오늘 여행 마치기」를 누르면 지나온 길과 들른 곳이 한 편의 여행기가 됩니다.',
                 '포스터 한 장으로 공유하고, 지나온 국도는 51선 기준으로 쌓입니다. 계정 없이 기기 안에 저장합니다.'],
        'shots': ['radar-t14.png', 'case-trip.png', 'share_card_8stops.png', 'my-t8.png'],
        'steps': [
            ['① 여행 마치기', '- 레이더의 「오늘 여행 마치기」'],
            ['② 여행기', '- 「동해 바닷길에서 생긴 일」 GPS 포스터 + 타임라인', f'- {CASE_LABEL}'],
            ['③ 공유 카드 (예시)', '- 지도 · 들른 곳 목록 · 「들른 발견 N」'],
            ['④ 마이', '- 여행기 목록 · 찜 · 국도 51선 수집 현황'],
        ],
    },
]

TOUR_APIS = [
    ('한국관광공사 국문 관광정보 서비스(KorService2) — 지역기반 관광정보 조회 areaBasedList2',
     ['시군구별 관광지·음식점·문화시설·쇼핑(5일장·상설시장) 목록 → 국도 51선 회랑 안의 장소 19,957건 적재',
      '→ 레이더·발견 카드의 후보']),
    ('국문 관광정보 — 공통정보 조회 detailCommon2 · 소개정보 조회 detailIntro2',
     ['개요(detailCommon2)와 영업시간·장날·전화(detailIntro2) → 신뢰도 점수(사진 35 · 영업시간 20 · 전화 15 · 주소 10 · 개요 10 · 추가사진 10, 60점 이상만 노출)와 카드 본문',
      '(예: 「3·8일에만 서는 장이라, 다음 장은 5일 뒤예요」)']),
    ('국문 관광정보 — 이미지정보 조회 detailImage2',
     ['관광공사 사진 → 발견 카드·장소 상세·「오늘 들른 곳」 자취에 원본 그대로(워터마크 포함)',
      '사진이 없는 장소는 카드로 보여주지 않음']),
    ('국문 관광정보 — 법정동코드 조회 ldongCode2',
     ['법정동 코드 → 시군구 순회 수집과 연관 관광지 이름 매칭의 지역 키',
      '(구 지역코드는 강원특별자치도 전환 뒤 누락이 있어 법정동 기준으로 조회)']),
    ('관광지별 연관 관광지 정보 서비스(TarRlteTarService1) — areaBasedList1 (baseYm 지정)',
     ['같은 방문자가 함께 찾는 관광지 순위 → 장소 상세 「함께 찾은 곳」과 노선 한 줄 「함께 찾는 곳이 많아요」',
      '두 기준월의 데이터가 있으면 순위 변화도 정렬에 참고(문구 「최근에도 함께 찾는 곳이 많아요」). 방문자 수·증가량·이동 순서·차량은 알 수 없어 말하지 않음']),
]

OTHER_DATA = [
    ('한국천문연구원 출몰시각 정보 OpenAPI (RiseSetInfoService)',
     ['전망 스팟이 있는 0.1° 격자 × 30일의 일몰 시각 캐시 → 일몰 60~20분 전 뷰포인트 카드·알림',
      '(예: 「곧 추암 촛대바위에 해가 져요」)']),
    ('소상공인시장진흥공단 전국전통시장표준데이터 (파일, CSV)',
     ['장날(3·8일 등) → 관광공사 시장 스팟에 결합 → 「오늘이 마침 장날」 카드와 다음 장날 안내',
      '사진·소개가 있는 관광공사 스팟에만 결합(신규 스팟 생성 없음)']),
    ('국토교통부 일반국도 도로중심선 (파일, SHP)',
     ['국도 51선 선형 → 홈 지도 폴리라인, 회랑 판정, 진입 지점 계산',
      '→ 여행기의 국도 51선 수집(지나온 점 맵매칭)']),
]

DIFF = [
    '• 목적지가 아니라 이동 중의 발견을 돕는다',
    '- 장소를 미리 정하는 대신, 현재 위치와 진행 방향에서 들를 후보를 제시. 길을 바꿔도 그대로 동작',
    '• 가까운 곳을 넘어 "지금 들를 이유"를 알려준다',
    '- 장날·일몰·이번 주에 끝나는 행사 등 시간 정보를 결합해, 같은 장소라도 시점에 따라 추천 의미를 달리함',
    '• 기존 내비와 역할을 나눈다',
    '- 길안내를 다시 만들지 않고 발견과 방문 기록에 집중. 「들르기」 선택 시 티맵·애플 지도가 안내하고, 하루가 여행기로 남음',
]
PLAN = [
    '• 핵심 가치 검증',
    '- 추천 후 「들르기」 선택률, 정차 감지 기반 방문 여부, 여행 후 질문으로 "계획에 없던 지역 방문"이 실제로 발생하는지 확인 (기기 내 집계·동의 기반)',
    '• 안전한 사용 상황 강화',
    '- 음성으로 카드에 응답(찜·들르기·넘기기)해 이동 중 화면 조작 없이 사용. 정차 감지 기준 정교화',
    '• 데이터 신뢰도 개선',
    '- 수집 파이프라인 상시화로 갱신 지연 축소, 임시휴무·정보 불확실 장소의 처리 기준 명시, 사용자 정차 흔적으로 "실제로 멈춘 곳" 보강',
    '• 시간 기반 발견 확대',
    '- 축제 기간 전체(현재는 이번 주에 끝나는 행사만)·야간 개장·계절 발견(단풍·벚꽃·해돋이)',
    '• 지역 확장·공유',
    '- 농산물 직판장·로컬푸드·지역 축제 데이터로 소도시 공백 보완, 여행기 공유 → 「남이 달린 길」(별점·댓글 없이 흔적만)',
]

# ──────────────────────────────── 채우기 ────────────────────────────────
prs = Presentation(str(SRC))
S = prs.slides

# 1 표지
t = find_table(S[0]).table
set_cell(t.cell(0, 1), TEAM, size=20, align=PP_ALIGN.CENTER)
set_cell(t.cell(1, 1), NAME, size=20, align=PP_ALIGN.CENTER)

# 2 서비스 소개
t = find_table(S[1]).table
set_cell(t.cell(0, 1), INTRO['name'], size=16)
set_cell(t.cell(1, 1), INTRO['type'], size=15)
# 링크 줄 — 하이퍼링크 런으로 (클릭되게)
tf = t.cell(1, 1).text_frame
p = tf.add_paragraph(); r = p.add_run(); r.text = 'App Store  '; r.font.size = Pt(12); r.font.color.rgb = INK
r = p.add_run(); r.text = APP_STORE_URL; r.font.size = Pt(12); r.hyperlink.address = APP_STORE_URL
if VIDEO_URL:
    p = tf.add_paragraph(); r = p.add_run(); r.text = '데모 영상  '; r.font.size = Pt(12); r.font.color.rgb = INK
    r = p.add_run(); r.text = VIDEO_URL; r.font.size = Pt(12); r.hyperlink.address = VIDEO_URL
set_cell(t.cell(2, 1), INTRO['summary'], size=13, bold_first=True)
set_cell(t.cell(3, 1), INTRO['why'], size=12, gap=5)

# 3 핵심기능 목록
t = find_table(S[2]).table
set_cell(t.cell(0, 1), FEATURES_LIST, size=12.5, gap=6)

# 4 이미지
t_sh = find_table(S[3]); t = t_sh.table
set_cell(t.cell(0, 1), '', size=12); set_cell(t.cell(1, 1), '', size=12)
left0 = Emu(t_sh.left).inches + Emu(t.columns[0].width).inches
top0 = Emu(t_sh.top).inches
r0, r1 = Emu(t.rows[0].height).inches, Emu(t.rows[1].height).inches
cw = Emu(t.columns[1].width).inches
if ICON.exists():
    # 양식은 '로고 또는 대표 이미지 1개' — 썸네일 한 장. 링크·문구는 2번 슬라이드에 있다.
    add_shot(S[3], ICON, left0, top0, cw, r0, pad=0.1)
else:
    missing.append(str(ICON))
detail = ['1_home.png', '2_depart.png', '3_card.png', '4_radar.png', '5_spot.png']
n = len(detail); gap = 0.18; ph = r1 - 0.2; pw = ph * 1320 / 2868
total = n * pw + (n - 1) * gap
x = left0 + (cw - total) / 2
for name in detail:
    p = img(name)
    if p: S[3].shapes.add_picture(str(p), Inches(x), Inches(top0 + r0 + 0.1), Inches(pw), Inches(ph))
    x += pw + gap

# 6 흐름도 — 원본(인덱스 5)을 채우고, 나머지는 복제해서 채운다
def fill_flow(slide, no, flow):
    remove_shape(slide, 19)  # 회색 안내 박스
    head = find_table(slide, '표 14').table
    set_cell(head.cell(0, 0), f'핵심 기능{no}', size=14, bold_first=True, align=PP_ALIGN.CENTER)
    set_cell(head.cell(0, 1), flow['title'], size=14, bold_first=True)
    set_cell(head.cell(1, 1), flow['desc'], size=11)
    body_sh = find_table(slide, '표 7'); body = body_sh.table
    left = Emu(body_sh.left).inches; top = Emu(body_sh.top).inches
    colw = [Emu(c.width).inches for c in body.columns]
    rowh = [Emu(r.height).inches for r in body.rows]
    y_img = top + rowh[0]
    x = left
    for i in range(4):
        set_cell(body.cell(1, i), '', size=10)
        set_cell(body.cell(2, i), flow['steps'][i], size=10.5, gap=2)
        shot = flow['shots'][i]
        name, crop = (shot, None) if isinstance(shot, str) else shot
        p = img(name)
        if p: add_shot(slide, p, x, y_img, colw[i], rowh[1], pad=0.08, crop=crop)
        x += colw[i]

fill_flow(S[5], 1, FLOWS[0])
extra = []
for i, flow in enumerate(FLOWS[1:], start=2):
    dst = duplicate_slide(prs, 5)
    fill_flow(dst, i, flow)
    extra.append(dst)

# 7 OpenAPI
t = find_table(S[6]).table
for i, (api, desc) in enumerate(TOUR_APIS):
    set_cell(t.cell(i * 2, 2), api, size=11, bold_first=True)
    set_cell(t.cell(i * 2 + 1, 2), desc, size=10)

# 8 기타 데이터
remove_shape(S[7], 6)
t = find_table(S[7]).table
for i, (nm, desc) in enumerate(OTHER_DATA):
    set_cell(t.cell(i * 2, 2), nm, size=11, bold_first=True)
    set_cell(t.cell(i * 2 + 1, 2), desc, size=10)

# 9 차별성·발전계획
t = find_table(S[8]).table
set_cell(t.cell(0, 1), DIFF, size=10.5, gap=2)
set_cell(t.cell(1, 1), PLAN, size=10.5, gap=2)

# 순서: 복제한 흐름도를 6번 슬라이드 뒤로. 5번(지역 특화)은 지운다 — 전국 서비스라 해당 없음.
count = len(prs.slides._sldIdLst)
for k, _ in enumerate(extra):
    move_slide(prs, count - len(extra) + k, 6 + k)   # 원본 흐름도(인덱스 5) 바로 뒤에 차례로
delete_slide(prs, 4)

OUT.parent.mkdir(exist_ok=True)
prs.save(str(OUT))
print('saved', OUT.relative_to(ROOT), '| slides:', len(prs.slides))
if missing: print('MISSING images:', missing)
